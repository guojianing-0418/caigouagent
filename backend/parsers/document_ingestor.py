"""统一文档摄取层。

把 PRD、规格书等文件转换成带来源元数据的文本单元。
第一版复用现有 openpyxl / PyMuPDF 路径，不引入新的重依赖。
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from ..models import DocumentIngestionDiagnostics, DocumentIngestionResult, DocumentTextUnit
from .excel_utils import cell_text


def ingest_document(path: str | Path | None, source_name: str = "文档", max_units: int = 800) -> DocumentIngestionResult:
    """摄取文档并返回文本单元与质量诊断。"""

    if not path:
        return DocumentIngestionResult(
            diagnostics=DocumentIngestionDiagnostics(
                parser="none",
            )
        )

    file_path = Path(path)
    if not file_path.exists():
        return DocumentIngestionResult(
            diagnostics=_diagnostics_with_warning(file_path.name, "missing", [f"文件不存在: {file_path}"])
        )

    suffix = file_path.suffix.lower()
    if suffix in {".xlsx", ".xlsm"}:
        return _ingest_excel(file_path, source_name=source_name, max_units=max_units)
    if suffix in {".txt", ".md", ".csv"}:
        return _ingest_text(file_path, source_name=source_name, max_units=max_units)
    if suffix == ".pdf":
        return _ingest_pdf(file_path, source_name=source_name, max_units=max_units)

    return DocumentIngestionResult(
        diagnostics=_diagnostics_with_warning(file_path.name, "unsupported", [f"暂不支持的规格书格式: {file_path.name}"])
    )


def _diagnostics_with_warning(file_name: str, parser: str, warnings: list[str]) -> DocumentIngestionDiagnostics:
    """创建带警告的诊断对象。"""

    return DocumentIngestionDiagnostics(
        file_name=file_name,
        parser=parser,
        warnings=warnings,
        suspicious_title_only=True,
    )


def _ingest_excel(path: Path, source_name: str, max_units: int) -> DocumentIngestionResult:
    """读取文档化 Excel，保留 sheet/行号/单元格数量诊断。"""

    units: list[DocumentTextUnit] = []
    warnings: list[str] = []
    non_empty_cell_count = 0
    parser = "openpyxl"

    wb = load_workbook(path, data_only=True, read_only=True)
    for ws in wb.worksheets:
        _repair_read_only_dimensions(ws, warnings)
        for row_number, row in enumerate(ws.iter_rows(values_only=True), start=1):
            values = [cell_text(value) for value in row]
            non_empty_values = [value for value in values if value]
            if not non_empty_values:
                continue
            non_empty_cell_count += len(non_empty_values)
            text = " | ".join(non_empty_values)
            units.append(
                DocumentTextUnit(
                    text=text,
                    source_name=source_name,
                    file_name=path.name,
                    sheet=ws.title,
                    row_number=row_number,
                    parser=parser,
                    raw_text=text,
                )
            )
            if len(units) >= max_units:
                warnings.append(f"文本单元超过 {max_units} 条，后续内容已截断。")
                return _build_result(path.name, parser, units, non_empty_cell_count, warnings)

    return _build_result(path.name, parser, units, non_empty_cell_count, warnings)


def _repair_read_only_dimensions(ws, warnings: list[str]) -> None:
    """修复部分 Excel 文件在只读模式下错误声明为 A1:A1 的维度。"""

    if not hasattr(ws, "reset_dimensions"):
        return
    try:
        dimension = ws.calculate_dimension()
    except Exception:
        return
    if dimension == "A1:A1":
        ws.reset_dimensions()
        warnings.append(f"工作表“{ws.title}”声明维度为 A1:A1，已重置维度后读取。")


def _ingest_text(path: Path, source_name: str, max_units: int) -> DocumentIngestionResult:
    """读取纯文本类文件。"""

    units: list[DocumentTextUnit] = []
    parser = "text"
    warnings: list[str] = []
    for row_number, line in enumerate(path.read_text(encoding="utf-8", errors="ignore").splitlines(), start=1):
        text = line.strip()
        if not text:
            continue
        units.append(
            DocumentTextUnit(
                text=text,
                source_name=source_name,
                file_name=path.name,
                row_number=row_number,
                parser=parser,
                raw_text=text,
            )
        )
        if len(units) >= max_units:
            warnings.append(f"文本单元超过 {max_units} 条，后续内容已截断。")
            break
    return _build_result(path.name, parser, units, len(units), warnings)


def _ingest_pdf(path: Path, source_name: str, max_units: int) -> DocumentIngestionResult:
    """抽取 PDF 可复制文本，主要用于规格书。"""

    try:
        import fitz  # PyMuPDF
    except Exception:
        return DocumentIngestionResult(
            diagnostics=_diagnostics_with_warning(path.name, "pymupdf", ["未安装 PyMuPDF，无法抽取 PDF 文本。请安装 requirements.txt 后重试。"])
        )

    units: list[DocumentTextUnit] = []
    parser = "pymupdf"
    warnings: list[str] = []
    with fitz.open(path) as doc:
        for page_index, page in enumerate(doc, start=1):
            text = page.get_text("text")
            for line in text.splitlines():
                clean = line.strip()
                if not clean:
                    continue
                units.append(
                    DocumentTextUnit(
                        text=clean,
                        source_name=source_name,
                        file_name=path.name,
                        page_number=page_index,
                        parser=parser,
                        raw_text=clean,
                    )
                )
                if len(units) >= max_units:
                    warnings.append(f"文本单元超过 {max_units} 条，后续内容已截断。")
                    return _build_result(path.name, parser, units, len(units), warnings)
    return _build_result(path.name, parser, units, len(units), warnings)


def _build_result(
    file_name: str,
    parser: str,
    units: list[DocumentTextUnit],
    non_empty_cell_count: int,
    warnings: list[str],
) -> DocumentIngestionResult:
    """汇总摄取结果与诊断。"""

    suspicious_title_only = len(units) <= 1
    if suspicious_title_only and units:
        warnings = [*warnings, "摄取结果只有 1 条文本单元，疑似只读取到标题或空文档。"]
    return DocumentIngestionResult(
        units=units,
        diagnostics=DocumentIngestionDiagnostics(
            file_name=file_name,
            parser=parser,
            text_unit_count=len(units),
            non_empty_cell_count=non_empty_cell_count,
            suspicious_title_only=suspicious_title_only,
            warnings=warnings,
        ),
    )
