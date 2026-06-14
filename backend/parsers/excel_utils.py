"""Excel 通用读取工具。

草 BOM、PRD、规格书和研发自提风险模板的 Excel 形态不完全一样，
因此这里提供两个低层工具：结构化行读取和文档化文本读取。
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook


def cell_text(value: Any) -> str:
    """把单元格值转成干净文本。"""

    if value is None:
        return ""
    return str(value).strip()


def read_sheet_rows(path: str | Path) -> dict[str, list[list[str]]]:
    """读取工作簿所有 sheet，返回二维文本数组。"""

    wb = load_workbook(path, data_only=True, read_only=True)
    result: dict[str, list[list[str]]] = {}
    for ws in wb.worksheets:
        rows: list[list[str]] = []
        for row in ws.iter_rows(values_only=True):
            values = [cell_text(v) for v in row]
            if any(values):
                rows.append(values)
        result[ws.title] = rows
    return result


def find_header_index(rows: list[list[str]], required_keyword: str) -> int:
    """寻找表头行。

    例如 BOM 表头里通常有“物料名称*”，研发自提模板里有“风险物料名称”。
    找不到时返回 0，保证后续逻辑可以继续给出友好错误或空结果。
    """

    for index, row in enumerate(rows[:20]):
        if any(required_keyword in cell for cell in row):
            return index
    return 0


def header_map(header_row: list[str]) -> dict[str, int]:
    """把表头名称映射为列号。"""

    return {name: index for index, name in enumerate(header_row) if name}


def get_by_header(row: list[str], headers: dict[str, int], names: list[str]) -> str:
    """按多个可能的表头名取值。

    例如“物料名称*”和“物料名称”都可以匹配。
    """

    for wanted in names:
        for header, index in headers.items():
            if wanted in header and index < len(row):
                return row[index].strip()
    return ""


def excel_as_text_lines(path: str | Path, max_lines: int = 500) -> list[str]:
    """把文档化 Excel 转成便于大模型或规则扫描的文本行。"""

    lines: list[str] = []
    for sheet_name, rows in read_sheet_rows(path).items():
        for row_number, row in enumerate(rows, start=1):
            text = " | ".join(cell for cell in row if cell)
            if text:
                lines.append(f"{sheet_name} R{row_number}: {text}")
            if len(lines) >= max_lines:
                return lines
    return lines

