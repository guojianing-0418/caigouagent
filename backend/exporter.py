"""风险物料 Excel 导出。"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import ensure_data_dirs, settings
from .models import DocumentSourceRef, RiskItem
from .risk_classification import ensure_risk_classification


BASE_EXPORT_HEADERS = ["风险物料名称", "所属模块", "风险类型", "风险原因", "来源与依据"]
CLASSIFICATION_HEADERS = [
    "风险确认方式",
    "主归口",
    "物料属性",
    "风险标签",
    "信息成熟度",
    "建议提问对象",
    "可发群问题",
    "需要补齐的信息",
    "分类依据",
]
EXPORT_HEADERS = BASE_EXPORT_HEADERS + CLASSIFICATION_HEADERS
QUESTION_DISTRIBUTION_HEADERS = [
    "风险物料名称",
    "风险类型",
    "主归口",
    "建议提问对象",
    "可发群问题",
    "需要补齐的信息",
    "信息成熟度",
    "来源与依据",
]
EVIDENCE_HEADERS = [
    "风险物料名称",
    "风险类型",
    "置信度",
    "待确认点",
    "来源名称",
    "文件名",
    "Sheet",
    "行号",
    "页码",
    "解析器",
    "证据原文",
]
HEADER_FILL = "2F6F5E"


def export_risks(project_id: str, project_name: str, risks: list[RiskItem]) -> Path:
    """导出最终风险物料清单。"""

    ensure_data_dirs()
    risks = ensure_risk_classification(risks)
    export_dir = settings.data_dir / "exports"
    export_dir.mkdir(parents=True, exist_ok=True)
    safe_name = "".join(ch if ch.isalnum() or ch in "-_." else "_" for ch in project_name)
    path = export_dir / f"{safe_name}_{project_id}_计划阶段风险物料.xlsx"

    wb = Workbook()
    ws = wb.active
    ws.title = "计划阶段风险物料"

    ws.append(EXPORT_HEADERS)
    _style_header(ws)

    main_risks = _sort_risks_for_main_sheet(risks)
    for risk in main_risks:
        ws.append(
            [
                risk.material_name,
                risk.module,
                risk.risk_type,
                risk.risk_reason,
                risk.source_basis,
                risk.risk_confirmation_method,
                risk.primary_owner,
                risk.material_attribute,
                _join_multi(risk.risk_tags),
                risk.information_maturity,
                risk.suggested_question_owner,
                _join_multi(risk.followup_questions),
                _join_multi(risk.missing_information),
                risk.classification_basis,
            ]
        )

    _style_body(ws, [24, 22, 20, 44, 80, 16, 14, 16, 36, 20, 24, 60, 50, 70])
    _merge_same_material_cells(ws)
    _append_question_distribution_sheet(wb, risks)
    _append_evidence_sheet(wb, risks)
    wb.save(path)
    return path


def create_rd_risk_template() -> Path:
    """生成研发自提风险 Excel 模板。"""

    ensure_data_dirs()
    path = settings.data_dir / "templates" / "研发自提风险模板.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "研发自提风险"
    ws.append(["风险物料名称", "原因"])
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 60
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2F6F5E")
        cell.alignment = Alignment(horizontal="center")
    wb.save(path)
    return path


def _append_question_distribution_sheet(wb: Workbook, risks: list[RiskItem]) -> None:
    """追加面向采购代表执行的问题分发清单。"""

    ws = wb.create_sheet("问题分发清单")
    ws.append(QUESTION_DISTRIBUTION_HEADERS)
    _style_header(ws)

    for risk in _sort_risks_for_main_sheet(risks):
        questions = risk.followup_questions or [""]
        for question in questions:
            ws.append(
                [
                    risk.material_name,
                    risk.risk_type,
                    risk.primary_owner,
                    risk.suggested_question_owner,
                    question,
                    _join_multi(risk.missing_information),
                    risk.information_maturity,
                    risk.source_basis,
                ]
            )

    _style_body(ws, [24, 20, 14, 24, 70, 50, 20, 80])


def _append_evidence_sheet(wb: Workbook, risks: list[RiskItem]) -> None:
    """追加结构化证据详情页，主表字段保持不变。"""

    ws = wb.create_sheet("证据详情")
    ws.append(EVIDENCE_HEADERS)
    _style_header(ws)

    for risk in risks:
        for row in _evidence_rows_for_risk(risk):
            ws.append(row)

    _style_body(ws, [24, 20, 12, 40, 18, 28, 18, 10, 10, 18, 90])


def _evidence_rows_for_risk(risk: RiskItem) -> list[list[str | int]]:
    """把单条风险转换为证据明细行；没有结构化证据时用来源依据兜底。"""

    confidence = _format_confidence(risk.confidence)
    unresolved = "\n".join(risk.unresolved_questions)
    evidence_items = risk.evidence_items or [
        DocumentSourceRef(
            source_name="来源与依据",
            excerpt=risk.source_basis,
        )
    ]
    rows: list[list[str | int]] = []
    for evidence in evidence_items:
        rows.append(
            [
                risk.material_name,
                risk.risk_type,
                confidence,
                unresolved,
                evidence.source_name,
                evidence.file_name,
                evidence.sheet,
                evidence.row_number or "",
                evidence.page_number or "",
                evidence.parser,
                evidence.excerpt or risk.source_basis,
            ]
        )
    return rows


def _format_confidence(confidence: float | None) -> str:
    """把 0-1 置信度显示成百分比。"""

    if confidence is None:
        return ""
    return f"{max(0, min(100, round(confidence * 100)))}%"


def _sort_risks_for_main_sheet(risks: list[RiskItem]) -> list[RiskItem]:
    """主表按物料名相邻展示，同物料内保持原顺序。"""

    indexed = list(enumerate(risks))
    indexed.sort(key=lambda item: (item[1].material_name or "\uffff", item[0]))
    return [risk for _, risk in indexed]


def _join_multi(values: list[str]) -> str:
    """把多值字段显示成 Excel 单元格中的换行文本。"""

    return "\n".join(str(item).strip() for item in values if str(item).strip())


def _merge_same_material_cells(ws: Worksheet) -> None:
    """合并主表第一列中连续相同的物料名。"""

    start_row = 2
    while start_row <= ws.max_row:
        material_name = ws.cell(row=start_row, column=1).value
        end_row = start_row
        while end_row + 1 <= ws.max_row and material_name and ws.cell(row=end_row + 1, column=1).value == material_name:
            end_row += 1
        if material_name and end_row > start_row:
            ws.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
            ws.cell(row=start_row, column=1).alignment = Alignment(vertical="center", wrap_text=True)
        start_row = end_row + 1


def _style_header(ws: Worksheet) -> None:
    """统一表头样式。"""

    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _style_body(ws: Worksheet, widths: list[int]) -> None:
    """统一列宽、换行和冻结首行。"""

    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    ws.freeze_panes = "A2"

